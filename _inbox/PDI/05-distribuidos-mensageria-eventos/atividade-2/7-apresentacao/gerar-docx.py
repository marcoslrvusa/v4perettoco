#!/usr/bin/env python3
"""Gera o DOCX desta atividade a partir de report.json (padrao PDI senior), incluindo as secoes de autoria."""
import json, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(HERE, "report.json"), encoding="utf-8") as f:
    D = json.load(f)

SKIP = {"Decisoes e tradeoffs", "Impacto no negocio", "Referencias de estudo"}

NEW = [
    ("Decisoes e tradeoffs", "ul", ["At-least-once do broker mais dedup no consumidor: entrega exactly-once observacional para o negocio, porque exactly-once de ponta a ponta nao existe em sistema distribuido.", "Outbox transacional: venda e evento gravados na mesma transacao na tabela outbox(event_id, payload, sent), entao falha na publicacao vira reprocessamento do outbox em vez de evento perdido. Custo: relay varrendo pendentes a cada 1s.", "Dedupe por idempotency_key antes de agir: o consumer consulta a tabela de processados antes de cobrar. Custo: store com TTL para nao encher.", "ACK so apos gravar a chave: a reentrega cai no dedupe, entao o mesmo evento 5x gera 1 cobranca. A chave combina event_id com identificador do negocio para nao colidir."]),
    ("Impacto no negocio", "p", ["Sem dedupe, reentregas geravam cobrancas duplicadas e cerca de 60h por mes de correcao manual. Com dedup a meta e zero duplicata e 100% dos handlers idempotentes, com correcao proxima de 0h por mes. O teste que injeta o mesmo evento 5x e afirma 1 cobranca da ao Financeiro previsibilidade: retry deixa de ser risco de debito duplo."]),
    ("Referencias de estudo", "ul", ["Curso: \"Event-Driven Architecture: From Theory to Practice\" (Udemy).", "Video: \"What is Idempotency?\" (YouTube, Hussein Nasser).", "Documento oficial: Apache Kafka Documentation, Exactly-once Semantics (kafka.apache.org).", "Documento oficial: PostgreSQL Documentation, INSERT ON CONFLICT (postgresql.org)."]),
]

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5); s.left_margin = Cm(3); s.right_margin = Cm(3)

def cover():
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(110)
    r = p.add_run('PDI'); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x8B,0x45,0x13)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(D['title']); r.bold = True; r.font.size = Pt(23)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(28)
    r = p.add_run('Documento Tecnico de PDI'); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x6B,0x7A,0x8A)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(16)
    for line in [f"Autor: {D['autor']}", f"Unidade: {D.get('unidade', 'FV Marketing / V4 Company')}", f"Data: {D.get('data', 'Agosto 2026')}",
                 f"Area: {D.get('area', '05 - Sistemas Distribuidos, Mensageria e Eventos')}", "Status: Entregue (desenvolvido)"]:
        rr = p.add_run(line + '\n'); rr.font.size = Pt(11); rr.font.color.rgb = RGBColor(0x49,0x55,0x60)
    doc.add_page_break()

def add_section(title, blocks):
    hh = doc.add_heading(title, level=2)
    for r in hh.runs: r.font.color.rgb = RGBColor(0x8B,0x1E,0x1E)
    for kind, *rest in blocks:
        if kind == "p":
            doc.add_paragraph(rest[0])
        elif kind in ("ul", "ol"):
            for x in rest[0]:
                doc.add_paragraph(x, style='List Bullet' if kind == "ul" else 'List Number')
        elif kind == "table":
            h, rows = rest
            t = doc.add_table(rows=1, cols=len(h)); t.style = 'Light Grid Accent 1'
            for i, c in enumerate(h): t.rows[0].cells[i].text = str(c)
            for r in rows:
                cells = t.add_row().cells
                for i, c in enumerate(r): cells[i].text = str(c)
        elif kind == "h":
            doc.add_heading(rest[0], level=3)
        elif kind == "note":
            p = doc.add_paragraph(); p.add_run("Nota: ").bold = True; p.add_run(rest[0])
        elif kind == "warn":
            p = doc.add_paragraph(); p.add_run("Atencao: ").bold = True; p.add_run(rest[0])
        elif kind == "code":
            p = doc.add_paragraph(rest[1]); p.style = doc.styles['No Spacing']
            for r in p.runs: r.font.name = 'Consolas'; r.font.size = Pt(9)

cover()
for title, blocks in D['sections']:
    if title.split(". ", 1)[-1] in SKIP:
        continue
    if isinstance(blocks, str):
        hh = doc.add_heading(title, level=2)
        for r in hh.runs: r.font.color.rgb = RGBColor(0x8B,0x1E,0x1E)
        continue
    add_section(title, blocks)
for title, kind, items in NEW:
    hh = doc.add_heading(title, level=2)
    for r in hh.runs: r.font.color.rgb = RGBColor(0x8B,0x1E,0x1E)
    if kind == "p":
        for x in items:
            doc.add_paragraph(x)
    else:
        for x in items:
            doc.add_paragraph(x, style='List Bullet')
out = os.path.join(HERE, "pdi-" + D['slug'] + ".docx")
doc.save(out)
print("DOCX:", out)
