#!/usr/bin/env python3
"""Gera o DOCX desta atividade a partir de report.json (padrao PDI senior), incluindo as secoes de autoria."""
import json, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(HERE, "report.json"), encoding="utf-8") as f:
    D = json.load(f)

SKIP = {"Decisoes e tradeoffs", "Impacto no negocio", "Referencias de estudo"}

NEW = [
    ("Decisoes e tradeoffs", "ul", ["RabbitMQ para filas de trabalho e Kafka para eventos de fluxo: a cobranca usa fila com 1 worker por mensagem e o Marketing e a Operacao assinam o mesmo topico em pub/sub. Troca um HTTP simples por um broker que precisa de operacao.", "ACK explicito com prefetch limitado: o consumer confirma depois de processar e recebe poucas mensagens por vez, entao um consumer lento nao estoura. Troca vazao por worker por estabilidade.", "DLQ apos N tentativas com runbook de reprocessamento: mensagem invalida vai para a DLQ em vez de sumir e volta pelo runbook-mensageria.md. Exige rotina de revisita em menos de 24h para a DLQ nao virar deposito esquecido.", "Publicacao fire-and-forget do venda.criada: o Vendas publica e segue em ~2ms sem esperar os outros times. Troca resposta imediata por consistencia eventual."]),
    ("Impacto no negocio", "p", ["Com alvo de 200 msg/s e zero perda em pico, o teste de 1k mensagens com o consumer derrubado prova que pico de campanha vira buffer no broker em vez de timeout em cascata. O P95 de ponta a ponta abaixo de 5s mantem Vendas, Financeiro e Marketing reagindo em segundos, o que reduz lead esquecido e retrabalho de conciliacao. O risco passa a ser operacional e conhecido: manter a DLQ revisitada em menos de 24h."]),
    ("Referencias de estudo", "ul", ["Curso: \"Apache Kafka Series: Learn Apache Kafka for Beginners\" (Udemy, Stephane Maarek).", "Video: \"RabbitMQ in 100 Seconds\" (YouTube, Fireship).", "Documento oficial: RabbitMQ Documentation (rabbitmq.com).", "Documento oficial: Apache Kafka Documentation (kafka.apache.org)."]),
]

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5); s.left_margin = Cm(3); s.right_margin = Cm(3)

def cover():
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(110)
    r = p.add_run('PDI'); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x8B,0x45,0x13)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(D['title']); r.bold = True; r.font.size = Pt(23)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(28)
    r = p.add_run('Documento Tecnico de PDI'); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x6B,0x7A,0x8A)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(16)
    for line in [f"Autor: {D['autor']}", f"Unidade: {D.get('unidade', 'FV Marketing / V4 Company')}", f"Data: {D.get('data', 'Agosto 2026')}",
                 f"Area: {D.get('area', '05 - Sistemas Distribuidos, Mensageria e Eventos')}", "Status: Entregue (desenvolvido)"]:
        rr = p.add_run(line + '\n'); rr.font.size = Pt(11); rr.font.color.rgb = RGBColor(0x49,0x55,0x60)
    doc.add_page_break()

def add_section(title, blocks):
    hh = doc.add_heading(title, level=2)
    for r in hh.runs: r.font.color.rgb = RGBColor(0x8B,0x1E,0x1E)
    for kind, *rest in blocks:
        if kind == "p":
            doc.add_paragraph(rest[0])
        elif kind in ("ul", "ol"):
            for x in rest[0]:
                doc.add_paragraph(x, style='List Bullet' if kind == "ul" else 'List Number')
        elif kind == "table":
            h, rows = rest
            t = doc.add_table(rows=1, cols=len(h)); t.style = 'Light Grid Accent 1'
            for i, c in enumerate(h): t.rows[0].cells[i].text = str(c)
            for r in rows:
                cells = t.add_row().cells
                for i, c in enumerate(r): cells[i].text = str(c)
        elif kind == "h":
            doc.add_heading(rest[0], level=3)
        elif kind == "note":
            p = doc.add_paragraph(); p.add_run("Nota: ").bold = True; p.add_run(rest[0])
        elif kind == "warn":
            p = doc.add_paragraph(); p.add_run("Atencao: ").bold = True; p.add_run(rest[0])
        elif kind == "code":
            p = doc.add_paragraph(rest[1]); p.style = doc.styles['No Spacing']
            for r in p.runs: r.font.name = 'Consolas'; r.font.size = Pt(9)

cover()
for title, blocks in D['sections']:
    if title.split(". ", 1)[-1] in SKIP:
        continue
    if isinstance(blocks, str):
        hh = doc.add_heading(title, level=2)
        for r in hh.runs: r.font.color.rgb = RGBColor(0x8B,0x1E,0x1E)
        continue
    add_section(title, blocks)
for title, kind, items in NEW:
    hh = doc.add_heading(title, level=2)
    for r in hh.runs: r.font.color.rgb = RGBColor(0x8B,0x1E,0x1E)
    if kind == "p":
        for x in items:
            doc.add_paragraph(x)
    else:
        for x in items:
            doc.add_paragraph(x, style='List Bullet')
out = os.path.join(HERE, "pdi-" + D['slug'] + ".docx")
doc.save(out)
print("DOCX:", out)
