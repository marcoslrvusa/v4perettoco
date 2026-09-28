#!/usr/bin/env python3
"""Gera o DOCX desta atividade a partir de report.json (padrao PDI senior), incluindo as secoes de autoria."""
import json, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(HERE, "report.json"), encoding="utf-8") as f:
    D = json.load(f)

SKIP = {"Decisoes e tradeoffs", "Impacto no negocio", "Referencias de estudo"}

NEW = [
    ("Decisoes e tradeoffs", "ul", ["consent_id obrigatorio no payload: so publica dado pessoal com base legal ativa; sem consentimento, o evento nao circula.", "subject_id em vez de CPF bruto no downstream: Vendas e Marketing operam com token e o Analytics recebe so hash sem reversao.", "AES em repouso mais TLS em transito: CPF e e-mail cifrados no broker, com campos sensiveis marcados com pii:true no schema lead_event.avsc.", "Retencao com TTL de 365 dias e purge por subject_id: o retention_purge.py varre e apaga, e o pedido de exclusao cai de 90 dias para menos de 1 dia, dentro do patamar de ate 15 dias."]),
    ("Impacto no negocio", "p", ["Com zero PII em texto puro e 100% dos fluxos com consentimento, o risco de autuacao pela ANPD e de dano de imagem cai porque o dado passa a ter rastro no mapa-dados.md e prazo definido. O apagamento em menos de 1 dia transforma o direito ao esquecimento em rotina operacional de uma varredura por subject_id, em vez de cacada manual por servico."]),
    ("Referencias de estudo", "ul", ["Curso: \"LGPD na Pratica\" (Udemy).", "Video: \"O que e a LGPD?\" (YouTube, SEBRAE).", "Documento oficial: Guia Orientativo da ANPD (gov.br/anpd).", "Documento oficial: Lei n. 13.709/2018 (planalto.gov.br)."]),
]

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5); s.left_margin = Cm(3); s.right_margin = Cm(3)

def cover():
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(110)
    r = p.add_run('PDI'); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x8B,0x45,0x13)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(D['title']); r.bold = True; r.font.size = Pt(23)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(28)
    r = p.add_run('Documento Tecnico de PDI'); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x6B,0x7A,0x8A)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(16)
    for line in [f"Autor: {D['autor']}", f"Unidade: {D.get('unidade', 'FV Marketing / V4 Company')}", f"Data: {D.get('data', 'Agosto 2026')}",
                 f"Area: {D.get('area', '05 - Sistemas Distribuidos, Mensageria e Eventos')}", "Status: Entregue (desenvolvido)"]:
        rr = p.add_run(line + '\n'); rr.font.size = Pt(11); rr.font.color.rgb = RGBColor(0x49,0x55,0x60)
    doc.add_page_break()

def add_section(title, blocks):
    hh = doc.add_heading(title, level=2)
    for r in hh.runs: r.font.color.rgb = RGBColor(0x8B,0x1E,0x1E)
    for kind, *rest in blocks:
        if kind == "p":
            doc.add_paragraph(rest[0])
        elif kind in ("ul", "ol"):
            for x in rest[0]:
                doc.add_paragraph(x, style='List Bullet' if kind == "ul" else 'List Number')
        elif kind == "table":
            h, rows = rest
            t = doc.add_table(rows=1, cols=len(h)); t.style = 'Light Grid Accent 1'
            for i, c in enumerate(h): t.rows[0].cells[i].text = str(c)
            for r in rows:
                cells = t.add_row().cells
                for i, c in enumerate(r): cells[i].text = str(c)
        elif kind == "h":
            doc.add_heading(rest[0], level=3)
        elif kind == "note":
            p = doc.add_paragraph(); p.add_run("Nota: ").bold = True; p.add_run(rest[0])
        elif kind == "warn":
            p = doc.add_paragraph(); p.add_run("Atencao: ").bold = True; p.add_run(rest[0])
        elif kind == "code":
            p = doc.add_paragraph(rest[1]); p.style = doc.styles['No Spacing']
            for r in p.runs: r.font.name = 'Consolas'; r.font.size = Pt(9)

cover()
for title, blocks in D['sections']:
    if title.split(". ", 1)[-1] in SKIP:
        continue
    if isinstance(blocks, str):
        hh = doc.add_heading(title, level=2)
        for r in hh.runs: r.font.color.rgb = RGBColor(0x8B,0x1E,0x1E)
        continue
    add_section(title, blocks)
for title, kind, items in NEW:
    hh = doc.add_heading(title, level=2)
    for r in hh.runs: r.font.color.rgb = RGBColor(0x8B,0x1E,0x1E)
    if kind == "p":
        for x in items:
            doc.add_paragraph(x)
    else:
        for x in items:
            doc.add_paragraph(x, style='List Bullet')
out = os.path.join(HERE, "pdi-" + D['slug'] + ".docx")
doc.save(out)
print("DOCX:", out)
