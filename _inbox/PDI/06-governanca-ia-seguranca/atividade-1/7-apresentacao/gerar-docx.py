#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o DOCX desta atividade (retrofit: tolerante a report.json + 3 novas secoes)."""
import json, os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(HERE, "report.json"), encoding="utf-8") as f:
    D = json.load(f)

TITLE = D.get("title", "Refatoracao de Modulo Legado com SOLID e Clean Architecture")
AUTOR = D.get("autor", "PDI Marcos Luciano - marcosluciano.rodrigues@v4company.com")
SLUG = D.get("slug", "praticas-engenharia-clean-code-a1")

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5); s.left_margin = Cm(3); s.right_margin = Cm(3)

def cover():
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(110)
    r = p.add_run("PDI"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x8B, 0x45, 0x13)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE); r.bold = True; r.font.size = Pt(23)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(28)
    r = p.add_run("Documento Tecnico de PDI"); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x6B, 0x7A, 0x8A)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(16)
    for line in [f"Autor: {AUTOR}", "Unidade: FV Marketing / V4 Company", "Data: Agosto 2026",
                 "Area: 02 Praticas de Engenharia e Clean Code", "Status: Entregue (desenvolvido)"]:
        rr = p.add_run(line + "\n"); rr.font.size = Pt(11); rr.font.color.rgb = RGBColor(0x49, 0x55, 0x60)
    doc.add_page_break()

def strip_html(s):
    s = re.sub(r"<br\s*/?>", "\n", s)
    s = re.sub(r"</(p|h[1-4]|li|tr|div)>", "\n", s)
    s = re.sub(r"<[^>]+>", "", s)
    s = ihtml_unescape(s)
    return re.sub(r"\n{3,}", "\n\n", s).strip()

def ihtml_unescape(s):
    import html as _h
    return _h.unescape(s)

def add_blocks(title, blocks):
    hh = doc.add_heading(title, level=2)
    for r in hh.runs:
        r.font.color.rgb = RGBColor(0x8B, 0x1E, 0x1E)
    if isinstance(blocks, str):
        txt = strip_html(blocks)
        for para in [x.strip() for x in txt.split("\n") if x.strip()][:120]:
            doc.add_paragraph(para[:1200])
        return
    for kind, *rest in blocks:
        if kind == "p":
            doc.add_paragraph(rest[0])
        elif kind in ("ul", "ol"):
            for x in rest[0]:
                doc.add_paragraph(x, style="List Bullet" if kind == "ul" else "List Number")
        elif kind == "table":
            h, rows = rest
            t = doc.add_table(rows=1, cols=len(h)); t.style = "Light Grid Accent 1"
            for i, c in enumerate(h):
                t.rows[0].cells[i].text = str(c)
            for r in rows:
                cells = t.add_row().cells
                for i, c in enumerate(r):
                    cells[i].text = str(c)
        elif kind == "h":
            doc.add_heading(rest[0], level=3)
        elif kind == "note":
            p = doc.add_paragraph(); p.add_run("Nota: ").bold = True; p.add_run(rest[0])
        elif kind == "warn":
            p = doc.add_paragraph(); p.add_run("Atencao: ").bold = True; p.add_run(rest[0])
        elif kind == "code":
            p = doc.add_paragraph(rest[1]); p.style = doc.styles["No Spacing"]
            for r in p.runs:
                r.font.name = "Consolas"; r.font.size = Pt(9)

cover()
for title, blocks in D.get("sections", []):
    add_blocks(title, blocks)
add_blocks("Decisoes e tradeoffs", [("ul", [
        "Clean Architecture escolhida sobre hexagonal puro e manter acoplado com E2E: testabilidade com ports desacoplados compensa o custo de mais arquivos, como registra o ADR-021.",
        "Dependencias de I/O como Protocol (LeadRepository, Notifier, Logger) com injecao no bootstrap: o servico passa a depender de abstracoes e o teste usa FakeRepo sem subir infra.",
        "Rollout em shadow por 1 sprint com reconciliacao diaria e corte em divergencia menor que 0,1 por cento, em vez de cutover direto: compara o modulo novo com o legado de 600 linhas sem expor listas de 5k a 80k leads a estado parcial.",
        "Meta de cobertura de 85 por cento nos testes de porta com mocks, em vez de teste manual: o legado tinha 0 por cento de cobertura e 4 dependencias de I/O acopladas, entao o gate quantitativo impede regressao silenciosa.",
        "Classes menores que 45 linhas com 1 responsabilidade por classe, aceitando mais arquivos: elimina SQL concatenado, falta de transacao e falha silenciosa de CRM que ja causou duplo contato."
    ])])
add_blocks("Impacto no negocio", [("p", "O modulo dispara 3 a 5 campanhas por dia para listas de 5k a 80k leads, entao cada falha silenciosa vira duplo contato e reclamacao real. Sair de 600 linhas com 0 por cento de cobertura para classes menores que 45 linhas com alvo de 85 por cento reduz o tempo de alteracao de deploy manual com teste manual para validacao automatica, baixa o risco de SQL injection e estado parcial sem transacao, e evita o custo de hotfix em base grande com rollback simples por feature flag.")])
add_blocks("Referencias de estudo", [("ul", [
        "Curso: Clean Architecture e SOLID com Python, na Alura.",
        "Video: SOLID em codigo Python na pratica, no YouTube.",
        "Doc oficial: Documentacao do Python sobre Protocol e tipagem estrutural, em docs.python.org.",
        "Doc oficial: Documentacao do pytest sobre fixtures e mocks, em docs.pytest.org."
    ])])

out = os.path.join(HERE, "pdi-" + SLUG + ".docx")
doc.save(out)
print("DOCX:", out)
