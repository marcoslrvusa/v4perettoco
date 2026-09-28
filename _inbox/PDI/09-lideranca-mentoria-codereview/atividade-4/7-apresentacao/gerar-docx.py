#!/usr/bin/env python3
"""Gera o DOCX da Atividade 4 da Trilha 09 (conteúdo inline, paleta PDI)."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
TITLE = "Rituais de Time Facilitados: Planning, Review e Retrospectiva"
SLUG = "lideranca-mentoria-codereview-a4"

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5); s.left_margin = Cm(3); s.right_margin = Cm(3)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(110)
r = p.add_run("PDI"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x8B, 0x45, 0x13)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(TITLE); r.bold = True; r.font.size = Pt(23)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(28)
r = p.add_run("Documento Técnico de PDI"); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x6B, 0x7A, 0x8A)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(16)
for line in ["Autor: PDI Marcos Luciano - marcosluciano.rodrigues@v4company.com",
             "Unidade: FV Marketing / V4 Company - Automação & Infraestrutura",
             "Data: Setembro 2026", "Área: Automação & Infraestrutura",
             "Status: Entregue (desenvolvido)"]:
    rr = p.add_run(line + "\n"); rr.font.size = Pt(11); rr.font.color.rgb = RGBColor(0x49, 0x55, 0x60)
doc.add_page_break()

def section(title, blocks):
    hh = doc.add_heading(title, level=2)
    for r in hh.runs: r.font.color.rgb = RGBColor(0x8B, 0x1E, 0x1E)
    for kind, *rest in blocks:
        if kind == "p":
            doc.add_paragraph(rest[0])
        elif kind in ("ul", "ol"):
            for x in rest[0]:
                doc.add_paragraph(x, style="List Bullet" if kind == "ul" else "List Number")
        elif kind == "table":
            h, rows = rest
            t = doc.add_table(rows=1, cols=len(h)); t.style = "Light Grid Accent 1"
            for i, c in enumerate(h): t.rows[0].cells[i].text = str(c)
            for row in rows:
                cells = t.add_row().cells
                for i, c in enumerate(row): cells[i].text = str(c)

section("1. Contexto", [("p", "Rituais existem no calendário e não funcionam: planning sem meta, review sem stakeholder, retro sem dono. Reunião sem facilitação consome hora cara e devolve zero decisão."),
    ("p", "A retro decorativa e o sintoma mais caro: o time aprende que falar não muda nada e para de falar.")])
section("2. Diagnóstico", [("p", "Só 10% dos rituais geram ata com dono, 25% das ações de retro são concluídas, 40% das sprints batem a meta e são 6h difusas por sprint. Causa raiz: ritual sem formato, sem timebox e sem definição de pronto."),
    ("table", ["Sintoma", "Atual"], [["Rituais com ata e dono", "10%"], ["Ações de retro concluídas", "25%"], ["Sprints na meta", "40%"], ["Tempo de rituais/sprint", "6h difusas"]])])
section("3. Solução", [("p", "Formato facilitado dos 3 rituais, roteiro de retro pronto, ata padrão em 24h e catálogo de 10 anti-patterns com correção."),
    ("ul", ["STD-01: planning 2h, review 45 min, retro 60 min", "STD-02: 10 anti-patterns com correção", "TPL roteiro de retro minuto a minuto", "TPL ata única dos 3 rituais"])])
section("4. Como funciona", [("ol", ["Planning: meta em 1 frase e itens com dono até a capacidade real", "Review: demo só do pronto, PO decide destino do feedback na hora", "Retro: coleta, votação, causa raiz do top 2, máximo 3 ações", "Ata em 24h e cobrança na planning seguinte", "Facilitação revezada, líder como membro"])])
section("5. Entregas", [("table", ["Arquivo", "Conteúdo"], [["STD-01-facilitacao-rituais.md", "Formato dos 3 rituais"], ["STD-02-anti-patterns.md", "10 anti-patterns"], ["TPL-roteiro-retro.md", "Roteiro de 60 min"], ["TPL-ata-rituais.md", "Ata única"]])])
section("6. Métricas de sucesso", [("table", ["Métrica", "Atual", "Meta"], [["Rituais com ata", "10%", "100% (meta)"], ["Ações concluídas", "25%", "80% (meta)"], ["Sprints na meta", "40%", "75% (meta)"], ["Tempo de rituais", "6h", "3h45 (meta)"]]),
    ("p", "Metas marcadas como (meta) por serem projetadas, não medidas. Rodar já na próxima sprint.")])

out = os.path.join(HERE, "pdi-" + SLUG + ".docx")
doc.save(out)
print("DOCX:", out)
