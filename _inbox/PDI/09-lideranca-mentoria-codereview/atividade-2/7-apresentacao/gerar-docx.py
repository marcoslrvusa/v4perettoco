#!/usr/bin/env python3
"""Gera o DOCX da Atividade 2 da Trilha 09 (conteudo inline, paleta PDI)."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
TITLE = "Mentoria e 1:1: Estrutura de Conversa e Plano de Evolucao"
SLUG = "lideranca-mentoria-codereview-a2"

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5); s.left_margin = Cm(3); s.right_margin = Cm(3)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(110)
r = p.add_run("PDI"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x8B, 0x45, 0x13)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(TITLE); r.bold = True; r.font.size = Pt(23)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(28)
r = p.add_run("Documento Tecnico de PDI"); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x6B, 0x7A, 0x8A)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(16)
for line in ["Autor: PDI Marcos Luciano - marcosluciano.rodrigues@v4company.com",
             "Unidade: FV Marketing / V4 Company - Automacao & Infraestrutura",
             "Data: Setembro 2026", "Area: Automacao & Infraestrutura",
             "Status: Entregue (desenvolvido)"]:
    rr = p.add_run(line + "\n"); rr.font.size = Pt(11); rr.font.color.rgb = RGBColor(0x49, 0x55, 0x60)
doc.add_page_break()

def section(title, blocks):
    hh = doc.add_heading(title, level=2)
    for r in hh.runs: r.font.color.rgb = RGBColor(0x8B, 0x1E, 0x1E)
    for kind, *rest in blocks:
        if kind == "p":
            doc.add_paragraph(rest[0])
        elif kind in ("ul", "ol"):
            for x in rest[0]:
                doc.add_paragraph(x, style="List Bullet" if kind == "ul" else "List Number")
        elif kind == "table":
            h, rows = rest
            t = doc.add_table(rows=1, cols=len(h)); t.style = "Light Grid Accent 1"
            for i, c in enumerate(h): t.rows[0].cells[i].text = str(c)
            for row in rows:
                cells = t.add_row().cells
                for i, c in enumerate(row): cells[i].text = str(c)

section("1. Contexto", [("p", "A 1:1 na operacao e status report com outro nome: lista de tasks, zero carreira, zero bloqueio real. Sem pauta fixa e sem registro, o crescimento depende de sorte e o burnout aparece na saida."),
    ("p", "Mentoria sem estrutura vira conselho solto que o liderado nao sabe aplicar na segunda-feira.")])
section("2. Diagnostico", [("p", "So 20% tem 1:1 recorrente, zero planos de evolucao ativos e media de 3 cancelamentos por liderado por mes. Causa raiz: conversa sem dono da pauta, sem roteiro e sem plano."),
    ("table", ["Sintoma", "Atual"], [["1:1 quinzenal recorrente", "20%"], ["Planos ativos", "0"], ["Cancelamentos/mes", "3 por liderado"]])])
section("3. Solucao", [("p", "Roteiro de 30 min em 4 blocos, banco de perguntas por tema e plano trimestral de ate 3 metas com evidencia observavel, prazo e apoio do lider."),
    ("ul", ["STD-01: formato, 4 blocos, perguntas, anti-patterns", "STD-02: metas, 70/20/10, notas 0/1/2", "TPL pauta de 1:1 por ciclo", "TPL plano trimestral do liderado"])])
section("4. Como funciona", [("ol", ["Check-in humano de 5 min", "Plano do trimestre de 10 min com evidencias", "Bloqueios e feedback bidirecional de 10 min", "Combinados de 5 min, maximo 3 com dono e prazo", "Fechamento trimestral de 30 min com notas"])])
section("5. Entregas", [("table", ["Arquivo", "Conteudo"], [["STD-01-estrutura-1-1.md", "Roteiro e perguntas"], ["STD-02-plano-de-evolucao.md", "Metas e acompanhamento"], ["TPL-pauta-1-1.md", "Ficha por ciclo"], ["TPL-plano-evolucao-liderado.md", "Ficha trimestral"]])])
section("6. Metricas de sucesso", [("table", ["Metrica", "Atual", "Meta"], [["1:1 recorrente", "20%", "100% (meta)"], ["Planos ativos", "0", "1 (meta)"], ["Combinados cumpridos", "sem medicao", "75% (meta)"], ["Cancelamentos/mes", "3", "0 (meta)"]]),
    ("p", "Metas marcadas como (meta) por serem projetadas, nao medidas. Revisao apos 6 semanas.")])

out = os.path.join(HERE, "pdi-" + SLUG + ".docx")
doc.save(out)
print("DOCX:", out)
