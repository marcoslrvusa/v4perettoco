#!/usr/bin/env python3
"""Gera o DOCX da Atividade 3 da Trilha 09 (conteudo inline, paleta PDI)."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
TITLE = "Delegacao e Feedback: Matriz de Delegacao e Feedback SBI"
SLUG = "lideranca-mentoria-codereview-a3"

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5); s.left_margin = Cm(3); s.right_margin = Cm(3)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(110)
r = p.add_run("PDI"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x8B, 0x45, 0x13)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(TITLE); r.bold = True; r.font.size = Pt(23)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(28)
r = p.add_run("Documento Tecnico de PDI"); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x6B, 0x7A, 0x8A)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(16)
for line in ["Autor: PDI Marcos Luciano - marcosluciano.rodrigues@v4company.com",
             "Unidade: FV Marketing / V4 Company - Automacao & Infraestrutura",
             "Data: Setembro 2026", "Area: Automacao & Infraestrutura",
             "Status: Entregue (desenvolvido)"]:
    rr = p.add_run(line + "\n"); rr.font.size = Pt(11); rr.font.color.rgb = RGBColor(0x49, 0x55, 0x60)
doc.add_page_break()

def section(title, blocks):
    hh = doc.add_heading(title, level=2)
    for r in hh.runs: r.font.color.rgb = RGBColor(0x8B, 0x1E, 0x1E)
    for kind, *rest in blocks:
        if kind == "p":
            doc.add_paragraph(rest[0])
        elif kind in ("ul", "ol"):
            for x in rest[0]:
                doc.add_paragraph(x, style="List Bullet" if kind == "ul" else "List Number")
        elif kind == "table":
            h, rows = rest
            t = doc.add_table(rows=1, cols=len(h)); t.style = "Light Grid Accent 1"
            for i, c in enumerate(h): t.rows[0].cells[i].text = str(c)
            for row in rows:
                cells = t.add_row().cells
                for i, c in enumerate(row): cells[i].text = str(c)

section("1. Contexto", [("p", "O lider centraliza por medo: delega a tarefa, segura a autoridade e refaz no final. O liderado ou espera ordem para tudo ou decide alem da conta."),
    ("p", "Feedback some por meses e explode generico, sem exemplo e sem pedido. Sem fato e sem saida, gera defesa, nao mudanca.")])
section("2. Diagnostico", [("p", "So 15% das tarefas delegadas tem criterio escrito, 2 SBIs por mes no time, 40% de retrabalho do lider e 25h semanais do lider em execucao. Causa raiz: delegacao binaria sem nivel e feedback sem metodo."),
    ("table", ["Sintoma", "Atual"], [["Delegadas com criterio", "15%"], ["SBIs por mes", "2"], ["Retrabalho do lider", "40%"], ["Horas operacionais do lider", "25h/sem"]])])
section("3. Solucao", [("p", "Matriz de 5 niveis, checklist de pacote delegavel, regra do indelegavel e roteiro SBI com exemplos prontos e controle mensal de proporcao reforco/correcao."),
    ("ul", ["STD-01: niveis, progressao, indelegavel, erros classicos", "STD-02: modelo SBI e exemplos prontos", "TPL matriz por liderado", "TPL roteiro SBI com controle mensal"])])
section("4. Como funciona", [("ol", ["Classificar o pacote na matriz de 1 a 5", "Escrever resultado, restricoes, nivel, checkpoint e pronto", "Executar sem interferencia ate o checkpoint", "Feedback SBI em ate 48h, 1 reforco para cada correcao", "Revisar o nivel a cada ciclo"])])
section("5. Entregas", [("table", ["Arquivo", "Conteudo"], [["STD-01-matriz-delegacao.md", "Niveis e progressao"], ["STD-02-feedback-sbi.md", "Modelo e exemplos"], ["TPL-matriz-delegacao.md", "Mapa por liderado"], ["TPL-roteiro-feedback-sbi.md", "Preparo e controle"]])])
section("6. Metricas de sucesso", [("table", ["Metrica", "Atual", "Meta"], [["Delegadas com criterio", "15%", "80% (meta)"], ["SBIs por mes", "2", "12 (meta)"], ["Retrabalho do lider", "40%", "10% (meta)"], ["Horas operacionais", "25h", "10h (meta)"]]),
    ("p", "Metas marcadas como (meta) por serem projetadas, nao medidas. Revisao de niveis ao fim do mes.")])

out = os.path.join(HERE, "pdi-" + SLUG + ".docx")
doc.save(out)
print("DOCX:", out)
