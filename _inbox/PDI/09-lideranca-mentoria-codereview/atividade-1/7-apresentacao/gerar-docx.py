#!/usr/bin/env python3
"""Gera o DOCX da Atividade 1 da Trilha 09 (conteudo inline, paleta PDI)."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
TITLE = "Code Review Efetivo: Checklist Objetivo, Cultura e Criterio de Barreira"
SLUG = "lideranca-mentoria-codereview-a1"

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5); s.left_margin = Cm(3); s.right_margin = Cm(3)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(110)
r = p.add_run("PDI"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x8B, 0x45, 0x13)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(TITLE); r.bold = True; r.font.size = Pt(23)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(28)
r = p.add_run("Documento Tecnico de PDI"); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x6B, 0x7A, 0x8A)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(16)
for line in ["Autor: PDI Marcos Luciano - marcosluciano.rodrigues@v4company.com",
             "Unidade: FV Marketing / V4 Company - Automacao & Infraestrutura",
             "Data: Setembro 2026", "Area: Automacao & Infraestrutura",
             "Status: Entregue (desenvolvido)"]:
    rr = p.add_run(line + "\n"); rr.font.size = Pt(11); rr.font.color.rgb = RGBColor(0x49, 0x55, 0x60)
doc.add_page_break()

def section(title, blocks):
    hh = doc.add_heading(title, level=2)
    for r in hh.runs: r.font.color.rgb = RGBColor(0x8B, 0x1E, 0x1E)
    for kind, *rest in blocks:
        if kind == "p":
            doc.add_paragraph(rest[0])
        elif kind in ("ul", "ol"):
            for x in rest[0]:
                doc.add_paragraph(x, style="List Bullet" if kind == "ul" else "List Number")
        elif kind == "table":
            h, rows = rest
            t = doc.add_table(rows=1, cols=len(h)); t.style = "Light Grid Accent 1"
            for i, c in enumerate(h): t.rows[0].cells[i].text = str(c)
            for row in rows:
                cells = t.add_row().cells
                for i, c in enumerate(row): cells[i].text = str(c)

section("1. Contexto", [("p", "O review na FV e loteria: PR gigante passa sem leitura e PR pequeno trava em gosto pessoal. Sem regra escrita, cada revisor inventa o criterio e o merge vira negociacao."),
    ("p", "O custo aparece depois do merge em hotfix, chamado de cliente e retrabalho que consome a sprint seguinte.")])
section("2. Diagnostico", [("p", "Tempo mediano de primeira resposta de 26h, 38% dos PRs acima de 400 linhas, 17% de retrabalho em 7 dias e 1,8 comentario acionavel por PR. Causa raiz: ausencia de definicao escrita do que barra o merge."),
    ("table", ["Sintoma", "Atual"], [["Primeira resposta", "26h"], ["PRs > 400 linhas", "38%"], ["Retrabalho pos merge", "17%"], ["Comentarios acionaveis/PR", "1,8"]])])
section("3. Solucao", [("p", "Politica de 7 barreiras objetivas, checklist de 6 blocos, SLA de 8h uteis (2h urgente), template de PR e CODEOWNERS com 2 approvals em area critica."),
    ("ul", ["STD-01: barreiras, SLA, teto de PR, papeis e vereditos", "STD-02: etiqueta de comentario e rito de impasse", "TPL checklist de review por categoria", "TPL descricao de PR com SLA"])])
section("4. Como funciona", [("ol", ["Autor abre PR com template e CI verde", "Revisor responde em 8h uteis (2h urgente)", "Checklist em 6 blocos, primeiro o que barra", "Veredito: approve, comment ou request changes", "Merge com squash e metricas semanais"])])
section("5. Entregas", [("table", ["Arquivo", "Conteudo"], [["STD-01-politica-code-review.md", "Barreiras, SLA, teto, papeis"], ["STD-02-cultura-e-etiqueta-review.md", "Comentarios e impasse"], ["TPL-checklist-review.md", "Checklist copiavel"], ["TPL-pr-descricao-e-sla.md", "Template de PR e SLA"]])])
section("6. Metricas de sucesso", [("table", ["Metrica", "Atual", "Meta"], [["Primeira resposta", "26h", "8h (meta)"], ["PRs > 400 linhas", "38%", "10% (meta)"], ["Retrabalho pos merge", "17%", "6% (meta)"], ["Comentarios acionaveis", "1,8", "3,0 (meta)"]]),
    ("p", "Metas marcadas como (meta) por serem projetadas, nao medidas. Revisao semanal pelo lider nos primeiros 2 meses.")])

out = os.path.join(HERE, "pdi-" + SLUG + ".docx")
doc.save(out)
print("DOCX:", out)
