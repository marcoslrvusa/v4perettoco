#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o DOCX desta atividade (retrofit: tolerante a report.json + 3 novas secoes)."""
import json, os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(HERE, "report.json"), encoding="utf-8") as f:
    D = json.load(f)

TITLE = D.get("title", "Mapeamento de Dominios com Domain-Driven Design (DDD)")
AUTOR = D.get("autor", "PDI Marcos Luciano - marcosluciano.rodrigues@v4company.com")
SLUG = D.get("slug", "praticas-engenharia-clean-code-a3")

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5); s.left_margin = Cm(3); s.right_margin = Cm(3)

def cover():
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(110)
    r = p.add_run("PDI"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x8B, 0x45, 0x13)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE); r.bold = True; r.font.size = Pt(23)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(28)
    r = p.add_run("Documento Tecnico de PDI"); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x6B, 0x7A, 0x8A)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(16)
    for line in [f"Autor: {AUTOR}", "Unidade: FV Marketing / V4 Company", "Data: Agosto 2026",
                 "Area: 02 Praticas de Engenharia e Clean Code", "Status: Entregue (desenvolvido)"]:
        rr = p.add_run(line + "\n"); rr.font.size = Pt(11); rr.font.color.rgb = RGBColor(0x49, 0x55, 0x60)
    doc.add_page_break()

def strip_html(s):
    s = re.sub(r"<br\s*/?>", "\n", s)
    s = re.sub(r"</(p|h[1-4]|li|tr|div)>", "\n", s)
    s = re.sub(r"<[^>]+>", "", s)
    s = ihtml_unescape(s)
    return re.sub(r"\n{3,}", "\n\n", s).strip()

def ihtml_unescape(s):
    import html as _h
    return _h.unescape(s)

def add_blocks(title, blocks):
    hh = doc.add_heading(title, level=2)
    for r in hh.runs:
        r.font.color.rgb = RGBColor(0x8B, 0x1E, 0x1E)
    if isinstance(blocks, str):
        txt = strip_html(blocks)
        for para in [x.strip() for x in txt.split("\n") if x.strip()][:120]:
            doc.add_paragraph(para[:1200])
        return
    for kind, *rest in blocks:
        if kind == "p":
            doc.add_paragraph(rest[0])
        elif kind in ("ul", "ol"):
            for x in rest[0]:
                doc.add_paragraph(x, style="List Bullet" if kind == "ul" else "List Number")
        elif kind == "table":
            h, rows = rest
            t = doc.add_table(rows=1, cols=len(h)); t.style = "Light Grid Accent 1"
            for i, c in enumerate(h):
                t.rows[0].cells[i].text = str(c)
            for r in rows:
                cells = t.add_row().cells
                for i, c in enumerate(r):
                    cells[i].text = str(c)
        elif kind == "h":
            doc.add_heading(rest[0], level=3)
        elif kind == "note":
            p = doc.add_paragraph(); p.add_run("Nota: ").bold = True; p.add_run(rest[0])
        elif kind == "warn":
            p = doc.add_paragraph(); p.add_run("Atencao: ").bold = True; p.add_run(rest[0])
        elif kind == "code":
            p = doc.add_paragraph(rest[1]); p.style = doc.styles["No Spacing"]
            for r in p.runs:
                r.font.name = "Consolas"; r.font.size = Pt(9)

cover()
for title, blocks in D.get("sections", []):
    add_blocks(title, blocks)
add_blocks("Decisoes e tradeoffs", [("ul", [
        "DDD explicito escolhido sobre schema unico: consistencia e linguagem comum compensam o custo de governanca, enquanto tabela unica acopla os 4 squads.",
        "Cada bounded context tem seu modelo e a integracao ocorre por eventos de dominio: evita que a entidade Contato continue com 3 modelos diferentes e que regra de um squad vaze para outro.",
        "So modelar o que tem regra de consistencia, com agregado de raiz unica e filhos sem identidade publica: previne over-engineering e concentra esforco onde ha invariante real.",
        "Validacao em workshop de linguagem ubiqua com Product e 2 squads contra 3 user stories antes de gerar schemas: garante que o mapa reflete o negocio e nao vira teoria.",
        "Anti-corruption layer na fronteira entre contextos: traduz o modelo vizinho sem poluir o contexto local, em vez de SQL cruzado entre dominios."
    ])])
add_blocks("Impacto no negocio", [("p", "Com 4 squads tocando Lead, Conta e agente sem vocabulario comum e 3 modelos diferentes de Contato, cada feature recria agregados e cada mudanca tem efeito cascata. Mapear 4 dominios com 0 modelos duplicados e ao menos 6 eventos definidos troca retrabalho e acoplamento oculto por fronteiras claras, reduz o tempo de impacto de alteracoes e evita o custo de carregar regra inconsistente para novas codificacoes.")])
add_blocks("Referencias de estudo", [("ul", [
        "Curso: Domain-Driven Design do zero, na Alura.",
        "Video: Bounded contexts e linguagem ubiqua na pratica, no YouTube.",
        "Doc oficial: Domain-Driven Design Reference, de Eric Evans, em domainlanguage.com.",
        "Doc oficial: Documentacao do Python sobre dataclasses para modelar agregados e eventos, em docs.python.org."
    ])])

out = os.path.join(HERE, "pdi-" + SLUG + ".docx")
doc.save(out)
print("DOCX:", out)
