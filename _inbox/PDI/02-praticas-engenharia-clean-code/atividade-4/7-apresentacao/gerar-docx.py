#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o DOCX desta atividade (retrofit: tolerante a report.json + 3 novas secoes)."""
import json, os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(HERE, "report.json"), encoding="utf-8") as f:
    D = json.load(f)

TITLE = D.get("title", "Estudo de Design Patterns aplicados ao ecossistema")
AUTOR = D.get("autor", "PDI Marcos Luciano - marcosluciano.rodrigues@v4company.com")
SLUG = D.get("slug", "praticas-engenharia-clean-code-a4")

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5); s.left_margin = Cm(3); s.right_margin = Cm(3)

def cover():
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(110)
    r = p.add_run("PDI"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x8B, 0x45, 0x13)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE); r.bold = True; r.font.size = Pt(23)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(28)
    r = p.add_run("Documento Tecnico de PDI"); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x6B, 0x7A, 0x8A)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(16)
    for line in [f"Autor: {AUTOR}", "Unidade: FV Marketing / V4 Company", "Data: Agosto 2026",
                 "Area: 02 Praticas de Engenharia e Clean Code", "Status: Entregue (desenvolvido)"]:
        rr = p.add_run(line + "\n"); rr.font.size = Pt(11); rr.font.color.rgb = RGBColor(0x49, 0x55, 0x60)
    doc.add_page_break()

def strip_html(s):
    s = re.sub(r"<br\s*/?>", "\n", s)
    s = re.sub(r"</(p|h[1-4]|li|tr|div)>", "\n", s)
    s = re.sub(r"<[^>]+>", "", s)
    s = ihtml_unescape(s)
    return re.sub(r"\n{3,}", "\n\n", s).strip()

def ihtml_unescape(s):
    import html as _h
    return _h.unescape(s)

def add_blocks(title, blocks):
    hh = doc.add_heading(title, level=2)
    for r in hh.runs:
        r.font.color.rgb = RGBColor(0x8B, 0x1E, 0x1E)
    if isinstance(blocks, str):
        txt = strip_html(blocks)
        for para in [x.strip() for x in txt.split("\n") if x.strip()][:120]:
            doc.add_paragraph(para[:1200])
        return
    for kind, *rest in blocks:
        if kind == "p":
            doc.add_paragraph(rest[0])
        elif kind in ("ul", "ol"):
            for x in rest[0]:
                doc.add_paragraph(x, style="List Bullet" if kind == "ul" else "List Number")
        elif kind == "table":
            h, rows = rest
            t = doc.add_table(rows=1, cols=len(h)); t.style = "Light Grid Accent 1"
            for i, c in enumerate(h):
                t.rows[0].cells[i].text = str(c)
            for r in rows:
                cells = t.add_row().cells
                for i, c in enumerate(r):
                    cells[i].text = str(c)
        elif kind == "h":
            doc.add_heading(rest[0], level=3)
        elif kind == "note":
            p = doc.add_paragraph(); p.add_run("Nota: ").bold = True; p.add_run(rest[0])
        elif kind == "warn":
            p = doc.add_paragraph(); p.add_run("Atencao: ").bold = True; p.add_run(rest[0])
        elif kind == "code":
            p = doc.add_paragraph(rest[1]); p.style = doc.styles["No Spacing"]
            for r in p.runs:
                r.font.name = "Consolas"; r.font.size = Pt(9)

cover()
for title, blocks in D.get("sections", []):
    add_blocks(title, blocks)
add_blocks("Decisoes e tradeoffs", [("ul", [
        "Adapter para CRM e LLM externos: isola o acoplamento a APIs de terceiro em tradutores atras de ports, entao novo fornecedor vira um adapter novo sem tocar o core.",
        "Strategy para roteamento de modelo de LLM e calculo variavel: troca if e else espalhados por estrategias selecionadas em runtime via factory selecionar.",
        "Observer para eventos de dominio com bus.assinar: reacoes como e-mail e estoque assinam eventos sem acoplar ao core, aceitando ordem nao garantida entre observadores.",
        "Singleton para clients rejeitado em favor de DI: evita estado global oculto e mantem o core testavel com fakes dos ports.",
        "Padrao aplicado so onde ha variacao real, com code review focado em valor: evita over-engineering e impede que padrao vire fim em si mesmo."
    ])])
add_blocks("Impacto no negocio", [("p", "Handlers de webhook e workers repetitivos com selecao de modelo por if e else fazem cada onboarding de fornecedor duplicar pontos de falha e travar o time. Com Adapter, Strategy e Observer mais piloto em ao menos 1 CRM, o onboarding cai para ate 2 dias atras de adapter testado, o teste do core usa fake sem chamar terceiro, e o custo de plugar parceiro novo deixa de ser reescrita do fluxo.")])
add_blocks("Referencias de estudo", [("ul", [
        "Curso: Design Patterns com Python, na Alura.",
        "Video: Strategy na pratica para trocar if else, no YouTube.",
        "Doc oficial: Catalogo de padroes com exemplos em Python, em refactoring.guru.",
        "Doc oficial: Documentacao do Python sobre abc e protocolos para ports e adapters, em docs.python.org."
    ])])

out = os.path.join(HERE, "pdi-" + SLUG + ".docx")
doc.save(out)
print("DOCX:", out)
