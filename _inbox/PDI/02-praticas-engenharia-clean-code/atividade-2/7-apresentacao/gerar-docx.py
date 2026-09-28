#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o DOCX desta atividade (retrofit: tolerante a report.json + 3 novas seções)."""
import json, os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(HERE, "report.json"), encoding="utf-8") as f:
    D = json.load(f)

TITLE = D.get("title", "Pipeline de Testes Automatizados (Unit/Integration/E2E) com 80% de Cobertura")
AUTOR = D.get("autor", "PDI Marcos Luciano - marcosluciano.rodrigues@v4company.com")
SLUG = D.get("slug", "praticas-engenharia-clean-code-a2")

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5); s.left_margin = Cm(3); s.right_margin = Cm(3)

def cover():
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(110)
    r = p.add_run("PDI"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x8B, 0x45, 0x13)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE); r.bold = True; r.font.size = Pt(23)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(28)
    r = p.add_run("Documento Técnico de PDI"); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x6B, 0x7A, 0x8A)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(16)
    for line in [f"Autor: {AUTOR}", "Unidade: FV Marketing / V4 Company", "Data: Agosto 2026",
                 "Área: 02 Práticas de Engenharia e Clean Code", "Status: Entregue (desenvolvido)"]:
        rr = p.add_run(line + "\n"); rr.font.size = Pt(11); rr.font.color.rgb = RGBColor(0x49, 0x55, 0x60)
    doc.add_page_break()

def strip_html(s):
    s = re.sub(r"<br\s*/?>", "\n", s)
    s = re.sub(r"</(p|h[1-4]|li|tr|div)>", "\n", s)
    s = re.sub(r"<[^>]+>", "", s)
    s = ihtml_unescape(s)
    return re.sub(r"\n{3,}", "\n\n", s).strip()

def ihtml_unescape(s):
    import html as _h
    return _h.unescape(s)

def add_blocks(title, blocks):
    hh = doc.add_heading(title, level=2)
    for r in hh.runs:
        r.font.color.rgb = RGBColor(0x8B, 0x1E, 0x1E)
    if isinstance(blocks, str):
        txt = strip_html(blocks)
        for para in [x.strip() for x in txt.split("\n") if x.strip()][:120]:
            doc.add_paragraph(para[:1200])
        return
    for kind, *rest in blocks:
        if kind == "p":
            doc.add_paragraph(rest[0])
        elif kind in ("ul", "ol"):
            for x in rest[0]:
                doc.add_paragraph(x, style="List Bullet" if kind == "ul" else "List Number")
        elif kind == "table":
            h, rows = rest
            t = doc.add_table(rows=1, cols=len(h)); t.style = "Light Grid Accent 1"
            for i, c in enumerate(h):
                t.rows[0].cells[i].text = str(c)
            for r in rows:
                cells = t.add_row().cells
                for i, c in enumerate(r):
                    cells[i].text = str(c)
        elif kind == "h":
            doc.add_heading(rest[0], level=3)
        elif kind == "note":
            p = doc.add_paragraph(); p.add_run("Nota: ").bold = True; p.add_run(rest[0])
        elif kind == "warn":
            p = doc.add_paragraph(); p.add_run("Atenção: ").bold = True; p.add_run(rest[0])
        elif kind == "code":
            p = doc.add_paragraph(rest[1]); p.style = doc.styles["No Spacing"]
            for r in p.runs:
                r.font.name = "Consolas"; r.font.size = Pt(9)

cover()
for title, blocks in D.get("sections", []):
    add_blocks(title, blocks)
add_blocks("Decisões e tradeoffs", [("ul", [
        "pytest com testcontainers e playwright escolhido sobre só unitários mockados: cobre 3 camadas com realismo e o custo maior de setup compensa, pois unitário sozinho e cego a integração.",
        "Unitário mira lógica pura com alvo de 90 por cento, integração mira ports com DB efêmero com alvo de 80 por cento, e E2E cobre só happy path: camadas rápidas seguram o merge e a camada lenta não trava o time.",
        "Gate de cobertura mínima de 80 por cento com --cov-fail-under=80 como required check no CI: impede piorar a cobertura sem perceber, saindo de 0 por cento e CI que só rodava lint.",
        "E2E em stage separado com retry, fora do caminho crítico do merge: evita que teste lento ou instável bloqueie o fluxo diário dos cerca de 30 módulos Python e nos JS.",
        "Fixtures isoladas com retry de 1 vez e isolamento contra estado global: sustenta tempo de unit e integração menor que 3 min e flaky rate menor que 1 por cento."
    ])])
add_blocks("Impacto no negócio", [("p", "O projeto tem cerca de 30 módulos Python mais nos JS e o CI atual só roda lint, então refactor de prompt ou tool injeta regressão em produção sem rede de segurança. O gate de 80 por cento com unit e integração em menos de 3 min troca dias de validação manual por minutos no CI, reduz regressão frequente para rara com flaky abaixo de 1 por cento, e evita o custo de corrigir defeito tarde, quando ele já chegou a main e a produção.")])
add_blocks("Referências de estudo", [("ul", [
        "Curso: Testes automatizados com pytest, na Alura.",
        "Vídeo: Piramide de testes na prática com Python, no YouTube.",
        "Doc oficial: Documentação do pytest sobre execução e cobertura, em docs.pytest.org.",
        "Doc oficial: Documentação do Coverage.py sobre medição com branch, em coverage.readthedocs.io."
    ])])

out = os.path.join(HERE, "pdi-" + SLUG + ".docx")
doc.save(out)
print("DOCX:", out)
