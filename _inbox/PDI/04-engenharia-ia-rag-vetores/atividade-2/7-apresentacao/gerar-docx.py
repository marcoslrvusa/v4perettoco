#!/usr/bin/env python3
"""Gera o DOCX desta atividade a partir de report.json (padrao PDI senior) + 3 secoes de autoria."""
import json, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(HERE, "report.json"), encoding="utf-8") as f:
    D = json.load(f)

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5); s.left_margin = Cm(3); s.right_margin = Cm(3)

def cover():
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(110)
    r = p.add_run('PDI'); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x8B,0x45,0x13)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(D['title']); r.bold = True; r.font.size = Pt(23)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(28)
    r = p.add_run('Documento Tecnico de PDI'); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x6B,0x7A,0x8A)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_before = Pt(16)
    for line in [f"Autor: {D['autor']}", f"Unidade: FV Marketing / V4 Company", f"Data: Agosto 2026",
                 f"Area: 04 Engenharia de IA, RAG e Vetores", "Status: Entregue (desenvolvido)"]:
        rr = p.add_run(line + '\n'); rr.font.size = Pt(11); rr.font.color.rgb = RGBColor(0x49,0x55,0x60)
    doc.add_page_break()

def add_section(title, blocks):
    hh = doc.add_heading(title, level=2)
    for r in hh.runs: r.font.color.rgb = RGBColor(0x8B,0x1E,0x1E)
    for kind, *rest in blocks:
        if kind == "p":
            doc.add_paragraph(rest[0])
        elif kind in ("ul", "ol"):
            for x in rest[0]:
                doc.add_paragraph(x, style='List Bullet' if kind == "ul" else 'List Number')
        elif kind == "table":
            h, rows = rest
            t = doc.add_table(rows=1, cols=len(h)); t.style = 'Light Grid Accent 1'
            for i, c in enumerate(h): t.rows[0].cells[i].text = str(c)
            for r in rows:
                cells = t.add_row().cells
                for i, c in enumerate(r): cells[i].text = str(c)
        elif kind == "h":
            doc.add_heading(rest[0], level=3)
        elif kind == "note":
            p = doc.add_paragraph(); p.add_run("Nota: ").bold = True; p.add_run(rest[0])
        elif kind == "warn":
            p = doc.add_paragraph(); p.add_run("Atencao: ").bold = True; p.add_run(rest[0])
        elif kind == "code":
            p = doc.add_paragraph(rest[1]); p.style = doc.styles['No Spacing']
            for r in p.runs: r.font.name = 'Consolas'; r.font.size = Pt(9)

BASE_SECTIONS = [('1. Contexto', [['p', 'Upgrade do baseline para hibrido BM25 mais vetorial com RRF e GraphRAG para relacoes do tipo cliente contrato fatura, onde similaridade pura falha em ID exato como CNPJ.']]), ('2. Diagnostico', [['p', 'Vetorial e otimo em sinonimo e ruim em ID exato. BM25 e otimo em exato e ruim em sinonimo. Relacao exige grafo.']]), ('3. Solucao', [['p', 'Retriever hibrido com RRF mais traversal no grafo, chunking semantico com overlap 128, e job noturno de rebuild incremental.']]), ('4. Entregas', [['ul', ['HYBRID-RAG.md', 'hybrid_rag.py e rag_hybrid.py com retrieve(q)', 'graph_schema.cypher e 001_rag_schema.sql']]]), ('5. Metricas', [['ul', ['hit@5 relacao maior ou igual a 0.9', 'hit@5 exato maior ou igual a 0.95', 'Precisao@5 de cerca de 42 por cento para cerca de 98 por cento', 'Latencia menor que 150 ms em 90 por cento das consultas', 'Meta precisao@5 maior ou igual a 95 por cento']]]), ('6. Status final', [['p', 'Desenvolvido e em homologacao. Aguarda revisao antes de producao.']])]

EXTRA_SECTIONS = [('9. Decisoes e tradeoffs', [['ul', ['BM25 mais vetorial com fusao RRF: aceitei complexidade extra para cobrir ID exato como CNPJ e sinonimo no mesmo retriever, porque cada metodo sozinho falha em um dos casos.', 'GraphRAG com traversal para relacoes cliente contrato fatura: assumi custo de rebuild incremental para responder perguntas relacionais que o vetorial nao resolve.', 'Chunking semantico com overlap 128: preservei contexto entre sentencas mesmo pagando mais tokens indexados.', 'Avaliacao em 30 perguntas (10 exatas, 10 sinonimos, 10 relacao) com hit@5: troquei teste informal por matriz que separa exato, sinonimo e relacao.', 'Meta de precisao@5 maior ou igual a 95 por cento e latencia menor que 150 ms com job noturno: equilibrei qualidade alta com atualizacao periodica do grafo.']]]), ('10. Impacto no negocio', [['p', 'O hibrido com hit@5 maior ou igual a 0.95 no exato e maior ou igual a 0.9 na relacao eleva a precisao@5 de cerca de 42 por cento para cerca de 98 por cento, o que reduz retrabalho de respostas vagas e viabiliza precificacao de busca relacional sem indexacao manual.']]), ('11. Referencias de estudo', [['ul', ['Curso: Advanced Retrieval for AI with Chroma, plataforma DeepLearning.AI.', 'Video: GraphRAG com busca vetorial mais grafo de conhecimento, plataforma YouTube, canal Microsoft Developer.', 'Doc oficial: Guia de BM25 e relevancia textual, documentacao oficial Elastic.', 'Doc oficial: Documentacao do pgvector com HNSW, documentacao oficial pgvector.']]])]

cover()
for title, blocks in (BASE_SECTIONS + EXTRA_SECTIONS):
    add_section(title, blocks)
out = os.path.join(HERE, "pdi-" + D['slug'] + ".docx")
doc.save(out)
print("DOCX:", out)
